import asyncio
import logging
from aiogram import Bot, Dispatcher, F
from aiogram.types import Message, FSInputFile
from aiogram.filters import CommandStart
from aiogram.fsm.storage.memory import MemoryStorage
import httpx
import json
from pptx import Presentation

BOT_TOKEN = "8957490808:AAEWFUdFyV8cpYE07rxbh-q2pHjsPrUXVYg"
API_KEY = "5911714ce3ffbc56f7064a9ad0708e0c" 
CHAT_API_URL = "https://api.kie.ai/gemini-3.1-pro/v1/chat/completions"

logging.basicConfig(level=logging.INFO)
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

@dp.message(CommandStart())
async def start(message: Message):
    await message.answer("Бот запущен. Напиши любую тему, и я создам презентацию.")

@dp.message(F.text)
async def generate(message: Message):
    topic = message.text
    msg = await message.answer("⏳ Генерирую...")
    
    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            prompt = f"Создай структуру из 3 слайдов на тему '{topic}'. JSON: [{'title': '...', 'content': '...'}, ...]"
            resp = await client.post(CHAT_API_URL, json={"messages": [{"role": "user", "content": prompt}]}, 
                                   headers={"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"})
            
            content = resp.json()['choices'][0]['message']['content'].replace("```json", "").replace("```", "").strip()
            data = json.loads(content)

            prs = Presentation()
            for item in data:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = item.get('title', '...')
                slide.placeholders[1].text = item.get('content', '')
            
            prs.save("final.pptx")
            await message.answer_document(FSInputFile("final.pptx"))
            await msg.delete()
    except Exception as e:
        await msg.edit_text(f"Ошибка: {str(e)}")

async def main():
    await bot.delete_webhook(drop_pending_updates=True)
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
